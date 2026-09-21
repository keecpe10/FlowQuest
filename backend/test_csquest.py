"""Integration checks with an isolated DB; never modifies classroom data."""
import os
import unittest
from unittest.mock import patch
import jwt
from uuid import uuid4
from sqlalchemy import create_engine, text
from sqlalchemy.engine import URL
from flask import Flask
from sqlalchemy import select, func
from app import db
from models import User, Role, Class, Course, CourseEnrollment, PointHistory, ActivityLog
from csquest.routes import csquest_bp
from csquest.models import Attempt, Progress, LearningWork, WorkSubmission
from csquest.curriculum import LEVELS, LEVEL_BY_ID

class CSQuestTests(unittest.TestCase):
    def setUp(self):
        self.env=patch.dict(os.environ, {'SECRET_KEY':'csquest-test-key-at-least-32-characters'})
        self.env.start()
        self.app=Flask(__name__)
        self.app.config.update(TESTING=True,SQLALCHEMY_DATABASE_URI='sqlite://',SQLALCHEMY_TRACK_MODIFICATIONS=False)
        self.admin_engine=None
        if os.environ.get('CSQUEST_TEST_POSTGRES'):
            url=URL.create('postgresql+psycopg2', username=os.environ['POSTGRES_USER'], password=os.environ['POSTGRES_PASSWORD'],host=os.environ['POSTGRES_HOST'],database=os.environ['POSTGRES_DB'])
            self.schema='csquest_test_'+uuid4().hex
            self.admin_engine=create_engine(url)
            with self.admin_engine.begin() as connection: connection.execute(text(f'CREATE SCHEMA {self.schema}'))
            self.app.config.update(SQLALCHEMY_DATABASE_URI=url, SQLALCHEMY_ENGINE_OPTIONS={'connect_args':{'options':f'-csearch_path={self.schema}'}})
        db.init_app(self.app)
        self.app.register_blueprint(csquest_bp)
        self.context=self.app.app_context();self.context.push()
        db.create_all()
        db.session.add_all([Role(role_id=1,role_name='teacher'),Role(role_id=2,role_name='student'),Class(class_id=1,class_name='ป.5/1'),Class(class_id=2,class_name='ป.5/2')])
        db.session.flush()
        for uid,role,room in [(1,1,None),(2,2,1),(3,2,2),(4,1,None),(5,2,None),(6,2,None)]:
            db.session.add(User(user_id=uid,username=f'user{uid}',first_name=f'Name{uid}',role_id=role,class_id=room,password_hash='unused',is_active=True,is_approved=True))
        db.session.flush()
        db.session.add_all([Course(course_id=1,course_name='Primary',teacher_id=1),Course(course_id=2,course_name='Other',teacher_id=4)])
        db.session.flush()
        db.session.add_all([CourseEnrollment(course_id=1,user_id=2),CourseEnrollment(course_id=2,user_id=3)])
        db.session.commit();self.client=self.app.test_client()

    def tearDown(self):
        db.session.remove();db.drop_all();db.engine.dispose();self.context.pop();self.env.stop()
        if self.admin_engine:
            with self.admin_engine.begin() as connection: connection.execute(text(f'DROP SCHEMA {self.schema} CASCADE'))
            self.admin_engine.dispose()

    def request(self,path,uid=2,method='GET',body=None):
        headers={} if uid is None else {'Authorization':'Bearer '+jwt.encode({'sub':str(uid)},os.environ['SECRET_KEY'],algorithm='HS256')}
        return self.client.open('/api/v1/csquest'+path,method=method,json=body,headers=headers)

    def complete(self,lid=1):
        self.assertEqual(self.request(f'/levels/{lid}/read',method='POST').status_code,200)
        result=self.request(f'/levels/{lid}/start',method='POST');self.assertEqual(result.status_code,200)
        aid=result.json['id']
        for _ in range(3):
            q=self.request(f'/attempts/{aid}/question').json['question']
            self.assertNotIn('answer',q)
            self.assertNotIn('explanation',q)
            source=next(v for v in LEVEL_BY_ID[lid]['questions'] if v['id']==q['id'])
            result=self.request(f'/attempts/{aid}/answer',method='POST',body={'question_id':q['id'],'choice':source['answer']})
            self.assertTrue(result.json['correct'])
        result=self.request(f'/attempts/{aid}/finish',method='POST');self.assertEqual(result.status_code,200)
        return aid,result.json

    def test_migration_preserves_accounts(self):
        import importlib.util
        from pathlib import Path
        from alembic.migration import MigrationContext
        from alembic.operations import Operations
        from sqlalchemy import inspect
        path=Path(__file__).parent/'migrations/versions/a819csquest_add_activities.py'
        spec=importlib.util.spec_from_file_location('cs_migration',path)
        migration=importlib.util.module_from_spec(spec);spec.loader.exec_module(migration)
        db.session.remove()
        with db.engine.begin() as connection:
            with Operations.context(MigrationContext.configure(connection)):
                migration.downgrade();migration.upgrade()
            self.assertEqual(connection.scalar(select(func.count()).select_from(User)),6)
            self.assertIn('cs_work_submission',inspect(connection).get_table_names())

    def test_all_curriculum_answers(self):
        from csquest.challenges import grade
        for level in LEVELS:
            for question in level['questions']:
                with self.subTest(level=level['id'],question=question['id']):
                    self.assertTrue(grade(question,question['answer'])['correct'])

    def test_identity_roles_and_scope(self):
        self.assertEqual(self.request('/me',uid=None).status_code,401)
        self.assertEqual(self.request('/session').json['user']['id'],2)
        self.assertEqual(self.request('/levels/1/read',method='POST',body=['invalid']).status_code,400)
        self.assertEqual(self.request('/teacher/students').status_code,403)
        self.assertEqual(self.request('/levels/1/start',uid=1,method='POST').status_code,403)
        self.assertEqual([u['id'] for u in self.request('/teacher/students',uid=1).json['students']],[2])
        self.assertEqual(self.request('/teacher/students/3/attempts',uid=1).status_code,404)
        self.assertEqual([u['id'] for u in self.request('/leaderboard',uid=5).json['entries']],[5])
        from io import BytesIO
        from openpyxl import load_workbook
        workbook=load_workbook(BytesIO(self.request('/teacher/export',uid=1).data))
        self.assertNotIn('user3',[row[1] for row in workbook.active.iter_rows(min_row=2,values_only=True)])
        workbook.close()
        db.session.get(User,2).is_active=False;db.session.commit()
        self.assertEqual(self.request('/me').status_code,401)

    def test_curriculum_unlock_and_awards(self):
        data=self.request('/curriculum').json
        self.assertEqual(len(data['levels']),51)
        for item in data['levels']:
            self.assertNotIn('questions',item)
        locked=next(l for l in LEVELS if l['prerequisite_id'])
        self.assertEqual(self.request(f"/levels/{locked['id']}/start",method='POST').status_code,403)
        aid,result=self.complete()
        self.assertEqual(result['profile']['score'],550)
        self.assertEqual(self.request(f'/attempts/{aid}/question',uid=3).status_code,404)
        self.request(f'/attempts/{aid}/finish',method='POST')
        self.complete()
        self.assertEqual(db.session.scalar(select(func.sum(PointHistory.points)).where(PointHistory.user_id==2)),550)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(PointHistory)),1)

    def test_xlsx_export_number_order_scores_and_scope(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.complete()
        db.session.get(User,2).student_number=10
        student=db.session.get(User,5);student.student_number=2;student.class_id=2
        student.username='0005';student.first_name='=HYPERLINK("https://example.com")'
        db.session.get(User,6).class_id=1
        db.session.get(User,3).student_number=1
        db.session.add_all([CourseEnrollment(course_id=1,user_id=5),CourseEnrollment(course_id=1,user_id=6)])
        db.session.commit()
        response=self.request('/teacher/export',uid=1)
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.mimetype,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('flowquest-grade5-scores.xlsx',response.headers['Content-Disposition'])
        workbook=load_workbook(BytesIO(response.data));sheet=workbook.active
        rows=list(sheet.iter_rows(min_row=2,values_only=True))
        self.assertEqual(sheet.cell(1,1).value,'เลขที่')
        self.assertEqual([row[0] for row in rows],[2,10,None])
        self.assertEqual([row[1] for row in rows],['0005','user2','user6'])
        self.assertEqual(rows[1][5],550);self.assertEqual(rows[1][10],550)
        self.assertEqual(sheet['B2'].data_type,'s');self.assertEqual(sheet['C2'].data_type,'s')
        self.assertEqual(sheet['A2'].data_type,'n');self.assertEqual(sheet['F3'].data_type,'n')
        self.assertEqual(sheet.freeze_panes,'D2');workbook.close()
        response=self.request('/teacher/export?classroom=ป.5/1',uid=1)
        workbook=load_workbook(BytesIO(response.data))
        self.assertEqual([r[1] for r in workbook.active.iter_rows(min_row=2,values_only=True)],['user2','user6']);workbook.close()
        workbook=load_workbook(BytesIO(self.request('/teacher/export?classroom=missing',uid=1).data))
        self.assertEqual(workbook.active.max_row,1);workbook.close()
        self.assertEqual(self.request('/teacher/export',uid=2).status_code,403)
        self.assertEqual(self.request('/teacher/export',uid=None).status_code,401)

    def test_teacher_reset_permissions(self):
        self.complete()
        path='/teacher/students/2/levels/1/reset'
        for uid,status in [(None,401),(2,403),(4,404)]:
            self.assertEqual(self.request(path,uid=uid,method='POST').status_code,status)
        self.assertEqual(self.request('/teacher/students/2/levels/9999/reset',uid=1,method='POST').status_code,404)
        self.assertEqual(self.request('/me').json['score'],550)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(ActivityLog)),0)

    def test_reset_one_level_preserves_other_data_and_can_replay(self):
        old_id,_=self.complete()
        next_level=next(l['id'] for l in LEVELS if l['prerequisite_id']==1)
        self.complete(next_level)
        work=LearningWork(user_id=2,kind='journal',scope_id=1,content={'learning':'keep'},state='submitted')
        db.session.add(work);db.session.flush()
        db.session.add(WorkSubmission(work_id=work.id,revision=1,content={'learning':'keep'}))
        db.session.add_all([Progress(user_id=3,level_id=1,completed=True,best_score=450),
                            PointHistory(user_id=2,source='mission',source_id=1,points=90),
                            PointHistory(user_id=3,source='csquest',source_id=1,points=450)])
        db.session.commit()
        path='/teacher/students/2/levels/1/reset'
        result=self.request(path,uid=1,method='POST')
        self.assertEqual(result.status_code,200,result.json)
        self.assertTrue(result.json['reset']);self.assertEqual(result.json['xp_removed'],550)
        self.assertEqual(result.json['student']['score'],550)
        self.assertEqual(result.json['student']['completed'],1)
        self.assertEqual(result.json['student']['attempts'],1)
        self.assertEqual(db.session.get(Attempt,old_id).status,'reset')
        p=db.session.scalar(select(Progress).where(Progress.user_id==2,Progress.level_id==1))
        self.assertTrue(p.read);self.assertFalse(p.completed);self.assertEqual(p.mushrooms,0)
        self.assertEqual(db.session.scalar(select(func.sum(PointHistory.points)).where(PointHistory.user_id==2)),640)
        self.assertEqual(db.session.scalar(select(func.sum(PointHistory.points)).where(PointHistory.user_id==3)),450)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(WorkSubmission)),1)
        self.assertEqual(db.session.get(LearningWork,work.id).content,{'learning':'keep'})
        log=db.session.scalar(select(ActivityLog));self.assertEqual(log.user_id,1);self.assertEqual(log.details['student_id'],2)
        self.assertFalse(self.request(path,uid=1,method='POST').json['reset'])
        self.assertEqual(db.session.scalar(select(func.count()).select_from(ActivityLog)),1)
        self.assertEqual(self.request(f'/attempts/{old_id}/finish',method='POST').status_code,409)
        self.complete()
        self.assertEqual(self.request('/me').json['score'],1100)
        self.assertEqual(db.session.scalar(select(func.sum(PointHistory.points)).where(PointHistory.user_id==2)),1190)

    def test_reset_stops_active_attempt(self):
        self.request('/levels/1/read',method='POST')
        aid=self.request('/levels/1/start',method='POST').json['id']
        result=self.request('/teacher/students/2/levels/1/reset',uid=1,method='POST')
        self.assertEqual(result.status_code,200);self.assertEqual(result.json['xp_removed'],0)
        for suffix in ('question','answer','finish','hit'):
            self.assertEqual(self.request(f'/attempts/{aid}/{suffix}',method='GET' if suffix=='question' else 'POST',body={} if suffix!='question' else None).status_code,409)
        self.assertEqual(self.request('/me').json['attempts'],0)
        fresh=self.request('/levels/1/start',method='POST')
        self.assertEqual(fresh.status_code,200);self.assertNotEqual(fresh.json['id'],aid)

    def test_concurrent_finish_and_reset(self):
        if not self.admin_engine:self.skipTest('Row locking requires PostgreSQL')
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        self.request('/levels/1/read',method='POST')
        aid=self.request('/levels/1/start',method='POST').json['id']
        for question in LEVEL_BY_ID[1]['questions']:
            self.request(f'/attempts/{aid}/answer',method='POST',body={'question_id':question['id'],'choice':question['answer']})
        db.session.remove()
        barrier=Barrier(2)
        def send(path,uid):
            with self.app.test_client() as client:
                token=jwt.encode({'sub':str(uid)},os.environ['SECRET_KEY'],algorithm='HS256')
                barrier.wait(timeout=5)
                return client.post('/api/v1/csquest'+path,headers={'Authorization':'Bearer '+token}).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            finish=pool.submit(send,f'/attempts/{aid}/finish',2)
            reset=pool.submit(send,'/teacher/students/2/levels/1/reset',1)
            self.assertIn(finish.result(timeout=15),(200,409));self.assertEqual(reset.result(timeout=15),200)
        self.assertEqual(self.request('/me').json['score'],0)
        self.assertEqual(db.session.get(Attempt,aid).status,'reset')
        self.assertEqual(db.session.scalar(select(func.coalesce(func.sum(PointHistory.points),0)).where(PointHistory.user_id==2)),0)

    def test_portfolio_scope_and_round_trip(self):
        self.complete()
        from csquest.portfolio import JOURNAL_FIELDS
        content={key:'สิ่งที่ได้เรียนรู้จากการทดลองและตรวจสอบคำตอบ' for key,_ in JOURNAL_FIELDS}
        saved=self.request('/portfolio/journal/1',method='PUT',body={'version':0,'content':content})
        self.assertEqual(saved.status_code,200,saved.json)
        version=saved.json['work']['version']
        submitted=self.request('/portfolio/journal/1/submit',method='POST',body={'version':version})
        self.assertEqual(submitted.status_code,200,submitted.json)
        entries=self.request('/teacher/portfolio',uid=1).json['entries'];self.assertEqual(len(entries),1)
        self.assertEqual(self.request('/teacher/portfolio',uid=4).json['entries'],[])
        entry=entries[0]
        body={'version':entry['version'],'submission_id':entry['submissions'][0]['id'],'feedback':'อธิบายได้ชัดเจน','outcome':'reviewed','rubric':{}}
        self.assertEqual(self.request(f"/teacher/portfolio/{entry['id']}/review",uid=4,method='POST',body=body).status_code,404)
        self.assertEqual(self.request(f"/teacher/portfolio/{entry['id']}/review",uid=1,method='POST',body=body).status_code,200)

if __name__=='__main__':unittest.main()
