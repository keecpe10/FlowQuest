"""Bounded XLSX parsing and validation for central student management."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from models import User

COLUMNS = [('เลขที่','student_number'),('ชื่อผู้ใช้','username'),('รหัสผ่าน','password'),
           ('ชื่อ','first_name'),('นามสกุล','last_name'),('ห้องเรียน','class_name'),
           ('ระดับชั้น','grade_level'),('ปีการศึกษา','academic_year'),('อีเมล','email')]
REQUIRED = {'student_number','username','password','first_name','last_name'}
MAX_ROWS = 1000

def student_number(value):
    if value is None or value == '': return None
    if isinstance(value,bool) or not re.fullmatch(r'[0-9]+',str(value).strip()):
        raise ValueError('เลขที่ต้องเป็นจำนวนเต็มบวก')
    result=int(value)
    if not 1 <= result <= 2147483647: raise ValueError('เลขที่ต้องเป็นจำนวนเต็มบวกไม่เกิน 2147483647')
    return result

def cell_text(cell):
    value=cell.value
    if value is None:return ''
    if cell.data_type=='f':raise ValueError('ไม่รองรับสูตร กรุณาใส่ค่าข้อมูลโดยตรง')
    if isinstance(value,bool):raise ValueError('รูปแบบข้อมูลไม่ถูกต้อง')
    if isinstance(value,(int,float)):
        if isinstance(value,float) and not value.is_integer():return str(value)
        number=int(value)
        # Preserve identifiers displayed with leading zeroes, e.g. 000123.
        if re.fullmatch(r'0{2,50}',cell.number_format or ''):
            return str(number).zfill(len(cell.number_format))
        return str(number)
    return str(value).strip()

def parse_students(file):
    if not file or not (file.filename or '').lower().endswith('.xlsx'):
        raise ValueError('กรุณาเลือกไฟล์ .xlsx')
    raw=file.read(5*1024*1024+1)
    if len(raw)>5*1024*1024:raise ValueError('ไฟล์ต้องมีขนาดไม่เกิน 5 MB')
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if len(archive.infolist())>500 or sum(i.file_size for i in archive.infolist())>30*1024*1024:
                raise ValueError('ไฟล์มีข้อมูลมากเกินไป กรุณาแบ่งไฟล์')
        workbook=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except ValueError:raise
    except Exception as error:raise ValueError('อ่านไฟล์ XLSX ไม่ได้ กรุณาตรวจสอบไฟล์') from error
    try:
        sheet=workbook.worksheets[0]
        # Do not trust workbook dimensions, which may be inflated or incorrect.
        sheet.reset_dimensions()
        iterator=sheet.iter_rows(max_col=32)
        header=next(iterator,None)
        if header is None:raise ValueError('ไม่พบหัวคอลัมน์ในชีตแรก')
        aliases={label:key for label,key in COLUMNS}|{key:key for _,key in COLUMNS}|{'ห้อง':'class_name'}
        columns={}
        for index,cell in enumerate(header):
            name=cell_text(cell)
            if name in aliases:
                key=aliases[name]
                if key in columns:raise ValueError('หัวคอลัมน์ซ้ำ: '+name)
                columns[key]=index
        missing=REQUIRED-set(columns)
        if missing:raise ValueError('ขาดคอลัมน์: '+', '.join(label for label,key in COLUMNS if key in missing))
        rows=[];errors=[]
        limits={'username':50,'password':128,'first_name':100,'last_name':100,'class_name':100,'grade_level':20,'academic_year':20,'email':120}
        for row_number,cells in enumerate(iterator,2):
            if row_number>MAX_ROWS+1:raise ValueError('รองรับไม่เกิน 1,000 แถวต่อไฟล์ กรุณาแบ่งไฟล์')
            if all(c.value is None for c in cells):continue
            row={key:'' for _,key in COLUMNS};row['row']=row_number
            try:
                for key,index in columns.items():row[key]=cell_text(cells[index])
                row['student_number']=student_number(row['student_number'])
                if row['student_number'] is None:raise ValueError('กรุณากรอกเลขที่')
                for key,label in [('username','ชื่อผู้ใช้'),('password','รหัสผ่าน'),('first_name','ชื่อ'),('last_name','นามสกุล')]:
                    if not row[key]:raise ValueError('กรุณากรอก'+label)
                if len(row['password'])<6:raise ValueError('รหัสผ่านต้องมีอย่างน้อย 6 ตัวอักษร')
                for key,limit in limits.items():
                    if len(row[key])>limit:raise ValueError(f'{next(label for label,k in COLUMNS if k==key)}ยาวเกิน {limit} ตัวอักษร')
                classroom=[row[key] for key in ('class_name','grade_level','academic_year')]
                if any(classroom) and not all(classroom):raise ValueError('กรอกห้องเรียน ระดับชั้น และปีการศึกษาให้ครบ หรือเว้นว่างทั้งหมด')
                if row['email'] and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',row['email']):raise ValueError('รูปแบบอีเมลไม่ถูกต้อง')
            except ValueError as error:errors.append({'row':row_number,'message':str(error)})
            rows.append(row)
        if not rows:raise ValueError('ไม่พบรายชื่อนักเรียนในชีตแรก')
        usernames=[r['username'] for r in rows];emails=[r['email'] for r in rows if r['email']]
        existing=User.query.filter((User.username.in_(usernames)) | (User.email.in_(emails))).all()
        used_names={u.username for u in existing};used_emails={u.email for u in existing if u.email}
        seen_names=set();seen_emails=set()
        for row in rows:
            for field,seen,used,label in [('username',seen_names,used_names,'ชื่อผู้ใช้'),('email',seen_emails,used_emails,'อีเมล')]:
                value=row[field]
                if not value:continue
                if value in used:errors.append({'row':row['row'],'message':label+'มีอยู่แล้วในระบบ'})
                elif value in seen:errors.append({'row':row['row'],'message':label+'ซ้ำในไฟล์'})
                seen.add(value)
        return rows,errors
    finally:workbook.close()

def template():
    workbook=Workbook();sheet=workbook.active;sheet.title='รายชื่อนักเรียน'
    sheet.append([label for label,_ in COLUMNS])
    # Blank template: no reusable passwords or accidental example accounts.
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='047857')
    from openpyxl.utils import get_column_letter
    for i in range(1,len(COLUMNS)+1):
        sheet.column_dimensions[get_column_letter(i)].width=22
        for row in range(2,102):sheet.cell(row,i).number_format='0' if i==1 else '@'
    sheet.freeze_panes='A2';sheet.auto_filter.ref='A1:I101'
    output=BytesIO();workbook.save(output);output.seek(0);return output
