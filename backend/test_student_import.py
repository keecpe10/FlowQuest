import unittest
from io import BytesIO
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, func, text
from werkzeug.security import check_password_hash
from app import db
from models import User, Class, ActivityLog
from student_routes import student_bp
from student_import import COLUMNS
import test_csquest

class StudentImportTests(unittest.TestCase):
    def setUp(self):
        self.fixture=test_csquest.CSQuestTests();self.fixture.setUp()
        self.app=self.fixture.app;self.app.register_blueprint(student_bp)
        db.session.get(User,1).is_super_admin=True;db.session.commit()
        if self.fixture.admin_engine:
            db.session.execute(text("SELECT setval(pg_get_serial_sequence('users','user_id'),6,true)"))
            db.session.execute(text("SELECT setval(pg_get_serial_sequence('classes','class_id'),2,true)"))
            db.session.commit()
        self.client=self.app.test_client()

    def tearDown(self):self.fixture.tearDown()

    def headers(self,uid=1):
        import os,jwt
        return {} if uid is None else {'Authorization':'Bearer '+jwt.encode({'sub':str(uid)},os.environ['SECRET_KEY'],algorithm='HS256')}

    def workbook(self,rows=None):
        wb=Workbook();ws=wb.active;ws.append([label for label,_ in COLUMNS])
        for row in (rows if rows is not None else [[1,'001234','secret12','เด็กชายหนึ่ง','ทดสอบ','ป.5/3','ป.5','2569','']]):ws.append(row)
        output=BytesIO();wb.save(output);output.seek(0);return output

    def upload(self,rows=None,uid=1,preview=False,file=None):
        return self.client.post('/api/v1/students/import'+('/preview' if preview else ''),
            data={'file':(file or self.workbook(rows),'students.xlsx')},headers=self.headers(uid))

    def test_preview_and_import_preserve_number_and_password(self):
        preview=self.upload(preview=True)
        self.assertEqual(preview.status_code,200,preview.json)
        self.assertEqual(preview.json['students'][0]['student_number'],1)
        self.assertNotIn('password',preview.json['students'][0]);self.assertNotIn('secret12',preview.text)
        self.assertIsNone(User.query.filter_by(username='001234').first())
        result=self.upload();self.assertEqual(result.status_code,201,result.json)
        user=User.query.filter_by(username='001234').one()
        self.assertEqual(user.student_number,1);self.assertTrue(check_password_hash(user.password_hash,'secret12'))
        self.assertEqual(user.school_class.class_name,'ป.5/3')
        self.assertEqual(self.upload().status_code,422)
        self.assertEqual(User.query.filter_by(username='001234').count(),1)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(ActivityLog)),1)

    def test_invalid_rows_are_atomic_and_report_row(self):
        rows=[[1,'new-one','secret12','หนึ่ง','ทดสอบ','New room','ป.5','2569',''],[0,'new-two','short','สอง','ทดสอบ','','','','']]
        response=self.upload(rows)
        self.assertEqual(response.status_code,422);self.assertEqual(response.json['errors'][0]['row'],3)
        self.assertEqual(User.query.filter(User.username.in_(['new-one','new-two'])).count(),0)
        self.assertEqual(Class.query.filter_by(class_name='New room').count(),0)
        rows[1][0]=2;rows[1][1]='new-one';rows[1][2]='secret12'
        response=self.upload(rows);self.assertEqual(response.status_code,422)
        self.assertIn('ซ้ำในไฟล์',str(response.json))

    def test_permissions_and_template(self):
        for uid,status in [(None,401),(2,403),(4,403)]:
            self.assertEqual(self.upload(uid=uid).status_code,status)
            self.assertEqual(self.client.get('/api/v1/students/import-template',headers=self.headers(uid)).status_code,status)
        response=self.client.get('/api/v1/students/import-template',headers=self.headers())
        self.assertEqual(response.status_code,200)
        wb=load_workbook(BytesIO(response.data));self.assertEqual(wb.active['A1'].value,'เลขที่');self.assertEqual(wb.active['B2'].number_format,'@');wb.close()

    def test_formula_corrupt_and_duplicate_accounts(self):
        row=[1,'=1+1','secret12','หนึ่ง','ทดสอบ','','','','']
        self.assertEqual(self.upload([row]).status_code,422)
        row[1]='user1';response=self.upload([row]);self.assertEqual(response.status_code,422)
        self.assertIn('มีอยู่แล้ว',str(response.json))
        self.assertEqual(self.upload(file=BytesIO(b'not xlsx')).status_code,400)
        self.assertEqual(self.upload([]).status_code,400)

    def test_numeric_identifier_leading_zero_and_limits(self):
        wb=Workbook();ws=wb.active;ws.append([label for label,_ in COLUMNS]);ws.append([2,123,'secret12','หนึ่ง','ทดสอบ','','','','']);ws['B2'].number_format='000000'
        output=BytesIO();wb.save(output);output.seek(0)
        self.assertEqual(self.upload(file=output).status_code,201)
        self.assertIsNotNone(User.query.filter_by(username='000123').first())
        rows=[[i+1,f'bulk{i}','secret12','หนึ่ง','ทดสอบ','','','',''] for i in range(1001)]
        self.assertEqual(self.upload(rows).status_code,400)
        self.assertEqual(User.query.filter(User.username.like('bulk%')).count(),0)

    def test_manual_create_update_number_and_clear(self):
        response=self.client.post('/api/v1/students/',json={'username':'manual','password':'secret12','student_number':'7'},headers=self.headers())
        self.assertEqual(response.status_code,201,response.json);uid=response.json['student']['user_id']
        self.assertEqual(response.json['student']['student_number'],7)
        for value in (-1,True,'1.5','2147483648'):
            self.assertEqual(self.client.patch(f'/api/v1/students/{uid}',json={'student_number':value},headers=self.headers()).status_code,400)
        response=self.client.patch(f'/api/v1/students/{uid}',json={'student_number':'8'},headers=self.headers())
        self.assertEqual(response.json['student']['student_number'],8)
        response=self.client.patch(f'/api/v1/students/{uid}',json={'student_number':''},headers=self.headers())
        self.assertIsNone(response.json['student']['student_number'])

    def test_student_number_migration_preserves_accounts(self):
        import importlib.util
        from pathlib import Path
        from alembic.migration import MigrationContext
        from alembic.operations import Operations
        path=Path(__file__).parent/'migrations/versions/c920_student_number.py'
        spec=importlib.util.spec_from_file_location('number_migration',path)
        migration=importlib.util.module_from_spec(spec);spec.loader.exec_module(migration)
        db.session.remove()
        with db.engine.begin() as connection:
            with Operations.context(MigrationContext.configure(connection)):
                migration.downgrade();migration.upgrade()
            self.assertEqual(connection.scalar(select(func.count()).select_from(User)),6)

if __name__=='__main__':unittest.main()
